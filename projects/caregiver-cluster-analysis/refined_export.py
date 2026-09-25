"""Build and independently recalculate the source-validated caregiver workbook.

This module never reads secrets. Its input is the validated source/scoring
bundle. All workbook publication is atomic and follows native recalculation.
"""
from __future__ import annotations

import argparse
from datetime import date, datetime, timezone
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
from typing import Any

import numpy as np
import pandas as pd

PROJECT_DIR = Path(__file__).resolve().parent
RULE_CODES = [f"R{i}" for i in range(1, 16)]
ACTION_ORDER = ["Pay now", "Check by hand", "Do not pay", "Incomplete - not eligible"]
EXCEL_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def excel_audit_spec() -> tuple[dict[str, str], list[str]]:
    """Independent native Excel logic; braces bind to visible audit inputs."""
    from refined_screening import DISPOSABLE_DOMAINS, SECTION_FLOORS_MIN, TOTAL_TIME_FLOOR_MIN, TFA_TIME_FLOOR_MIN
    fields = ["family_information_form_complete", "values_complete", "tfa_complete", "demographics_complete"]
    completed = "AND(" + ",".join("{" + f + "}=2" for f in fields) + ")"
    timing_names = ["Family Information minutes", "Values minutes", "Thoughts and feelings minutes", "Demographics minutes"]
    low_sections = [f"AND(ISNUMBER({{{name}}}),{{{name}}}<{floor})" for name,floor in zip(timing_names,SECTION_FLOORS_MIN.values())]
    total = "+".join("{"+name+"}" for name in timing_names)
    numeric_times = ",".join(f"ISNUMBER({{{name}}})" for name in timing_names)
    finite_children = "AND(ISNUMBER({Children maximum}),{Children count is open ended}=0)"
    impossible = f"AND(ISNUMBER({{Autistic children minimum}}),{finite_children},{{Autistic children minimum}}>{{Children maximum}})"
    conditions = {
        "R1":f"AND({numeric_times},{total}<{TOTAL_TIME_FLOOR_MIN})",
        "R2":f"AND(ISNUMBER({{Thoughts and feelings minutes}}),{{Thoughts and feelings minutes}}<{TFA_TIME_FLOOR_MIN})",
        "R3":"OR("+",".join(low_sections)+")",
        "R4":"AND(ISNUMBER({Minimum rating block SD}),{Minimum rating block SD}<0.15)",
        "R5":"AND({Rating items available}>=20,{Rating fraction answered}>=0.8,{Records sharing response fingerprint}>=2)",
        "R6":"AND(ISNUMBER({Start hour}),{Other sign-ups within two minutes}>=2)",
        "R7":"AND(ISNUMBER({Maximum narrative cosine similarity}),{Maximum narrative cosine similarity}>0.9)",
        "R8":f"OR(AND(ISNUMBER({{Autistic children minimum}}),{{Autistic children minimum}}=0,{{Autism follow-up answered}}=1),AND({finite_children},{{Selected child age bands}}>{{Children maximum}}),{impossible})",
        "R9":'OR(AND(ISNUMBER({Reported caregiver age}),OR({Reported caregiver age}<18,{Reported caregiver age}>100)),{ZIP format invalid}=1,AND(ISNUMBER({Parent age at first birth lower}),ISNUMBER({Parent age at first birth upper}),OR({Parent age at first birth upper}<10,{Parent age at first birth lower}>60)))',
        "R10":"OR("+",".join('{Email domain}="'+d+'"' for d in sorted(DISPOSABLE_DOMAINS))+")",
        "R11":f'AND({completed},{{Email address}}="")',
        "R12":"AND(ISNUMBER({Start hour}),{Start hour}>=0,{Start hour}<5)",
        "R13":"AND({Attention check available}=1,{Attention check answered}=1,OR({Band-Aid selected}<>1,{Other attention options selected}>0))",
        "R14":"OR({She and man pairing}=1,{He and woman pairing}=1)",
        "R15":"AND({Explicit no disability response}=1,ISNUMBER({Autistic children minimum}),{Autistic children minimum}>=1)",
    }
    flagged = lambda rule: "{"+rule+' result}="FLAG"'
    timing = f'IF(OR({flagged("R1")},{flagged("R2")}),1,0)'
    copying = f'IF(OR(AND({flagged("R5")},{{Rating pattern has at least three values}}=1),{flagged("R7")}),1,0)'
    logic = f'IF({impossible},1,0)'
    attention = f'IF(AND({flagged("R13")},{{Attention accessibility or opt-out response}}=0),1,0)'
    families = "+".join([timing,copying,logic,attention])
    source_issue = 'OR({Source mapping issue}=1,{Source project verified}<>"Yes")'
    review = "OR("+",".join(flagged(r) for r in ["R1","R2","R5","R7","R8","R9","R13","R15"])+f",{source_issue},{{Timing data issue}}=1,AND({completed},{{Sections with a recorded time}}=0),AND({completed},{{Attention check available}}=1,{{Attention check answered}}=0))"
    contact = f'OR({{Email usable}}<>"Yes",{flagged("R10")},{{Email confirmation differs}}=1,{{Responses sharing payment email}}>1)'
    conditions["decision"] = f'IF(NOT({completed}),"Incomplete - not eligible",IF({source_issue},"Check by hand",IF(({families})>=2,"Do not pay",IF(OR({review},{contact}),"Check by hand","Pay now"))))'
    evidence = list(dict.fromkeys(fields + timing_names + [
        "Minimum rating block SD","Rating items available","Rating fraction answered","Records sharing response fingerprint",
        "Rating pattern has at least three values","Start hour","Other sign-ups within two minutes","Maximum narrative cosine similarity",
        "Autistic children minimum","Children maximum","Children count is open ended","Autism follow-up answered","Selected child age bands",
        "Reported caregiver age","US residence selected","ZIP provided","ZIP format invalid","Parent age at first birth lower","Parent age at first birth upper",
        "Email domain","Email address","Email usable","Email confirmation differs","Responses sharing payment email",
        "Attention check available","Attention check answered","Band-Aid selected","Other attention options selected","Attention accessibility or opt-out response",
        "She and man pairing","He and woman pairing","Pronoun response","Gender response","Explicit no disability response",
        "Source mapping issue","Source project verified","Timing data issue","Sections with a recorded time"] ))
    return conditions,evidence


def excel_column(index: int) -> str:
    result = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _clean(value: Any) -> Any:
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, (np.integer, int)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, (pd.Timestamp, datetime, date)):
        stamp = pd.Timestamp(value)
        if stamp.tzinfo is not None:
            stamp = stamp.tz_localize(None)
        return {"excelDate": stamp.isoformat() + "Z"}
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, default=str)
    text = str(value)
    if text in {"nan", "NaN", "<NA>", "None", "NaT"}:
        return None
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _width(label: str) -> float:
    if label == "Study":
        return 66
    if label == "Rule":
        return 12
    if label in {"Points", "Evidence family"}:
        return 18
    if label == "Check":
        return 48
    if label in {"Definition", "Live REDCap project title"}:
        return 72
    if label in {"PID", "Record ID", "REDCap PID"}:
        return 13
    if label in {"Recommended action", "Action from Excel", "Score-only proposed action"}:
        return 29
    if any(word in label.lower() for word in ("reason", "why", "evidence", "description", "notes", "plan")):
        return 74
    if "email" in label.lower():
        return 34
    if "parity" in label or "result" in label:
        return 18
    return min(40, max(18, len(label) * 0.60))


def _frame_sheet(name: str, frame: pd.DataFrame, **kwargs: Any) -> dict:
    frame = frame.copy()
    if frame.columns.duplicated().any():
        raise ValueError(f"Duplicate workbook columns in {name}")
    if not len(frame.columns):
        frame = pd.DataFrame(columns=["Result"])
    headers = [str(c) for c in frame.columns]
    formats = []
    for label in headers:
        if label in {"PID", "REDCap PID", "Record ID", "Response key"}:
            formats.append("@")
        elif "minutes" in label.lower() or "similarity" in label.lower() or "deviation" in label.lower():
            formats.append("0.00")
        elif pd.api.types.is_datetime64_any_dtype(frame[label]):
            formats.append("mm/dd/yy hh:mm")
        else:
            formats.append(None)
    return {"name": name, "headers": headers,
            "rows": [[_clean(v) for v in row] for row in frame.itertuples(index=False, name=None)],
            "widths": [_width(c) for c in headers], "formats": formats, **kwargs}


def _public_records(scored: pd.DataFrame, rules: dict) -> pd.DataFrame:
    frame = scored.rename(columns={"REDCap PID": "PID"}).copy()
    required = ["Study", "PID", "Record ID", "Recommended action", "Why this action", "Risk score", "Survey finished"]
    if set(required) - set(frame):
        raise ValueError(f"Screening output missing columns: {sorted(set(required) - set(frame))}")
    preferred = required[:5] + ["Risk score", "Survey finished", "Review priority", "Review plan", "Final plan", "Bot assessment", "Independent evidence families",
        "Checks flagged", "Serious checks flagged (prompt weights)", "Mild checks flagged (prompt weights)", "Which checks were flagged", "Points from each check",
        "Email address", "University email address", "Email usable", "Started", "Last section handed in",
        "Family Information minutes", "Values minutes", "Thoughts and feelings minutes", "Demographics minutes", "Whole survey minutes",
        "Sections with a recorded time", "Minutes from sign-up to last section", "Other sign-ups within two minutes",
        "Maximum narrative cosine similarity", "Minimum rating block SD", "Rating fraction answered", "Records sharing response fingerprint",
        "Rules evaluated", "Source project verified", "Source mapping issue", "Source mapping notes", "Recruitment population", "Reference cohort designation",
        "Score-only proposed action", "Response key"]
    out = frame[[c for c in preferred if c in frame]].copy()
    for c in ["Started", "Last section handed in"]:
        if c in out:
            out[c] = pd.to_datetime(out[c], errors="coerce", format="mixed")
    out["PID"] = out["PID"].astype(str)
    out["Record ID"] = out["Record ID"].astype(str)
    for code in RULE_CODES:
        out[f"{code}: {rules[code]['label']}"] = frame[f"{code} status"]
    return out


def _rule_audit(scored: pd.DataFrame, rules: dict, conditions: dict[str, str], evidence: list[str]) -> dict:
    """Recompute rules from visible normalized inputs; audit never feeds outputs."""
    audit = scored.rename(columns={"REDCap PID": "PID"})[["Study", "PID", "Record ID"]].copy()
    audit["PID"], audit["Record ID"] = audit["PID"].astype(str), audit["Record ID"].astype(str)
    for c in ["Score parity", "Decision parity", *[f"{r} parity" for r in RULE_CODES],
              "Score from Excel", "Action from Excel", *[f"{r} result" for r in RULE_CODES]]:
        audit[c] = None
    audit["Python risk score"] = scored["Risk score"].to_numpy()
    audit["Python recommended action"] = scored["Recommended action"].to_numpy()
    for c in evidence:
        if c not in scored:
            raise ValueError(f"Missing rule-audit evidence: {c}")
        if c not in audit:
            audit[c] = scored[c].to_numpy()
    audit = audit.copy()
    for r in RULE_CODES:
        audit[f"{r} Python flag"] = scored[r].astype(int).to_numpy()
        audit[f"{r} availability"] = scored[f"{r} status"].to_numpy()
        audit[f"{r} evidence"] = scored[f"{r} evidence"].to_numpy()
    spec = _frame_sheet("Rule Audit", audit)
    spec["formulas"], spec["expected"] = {}, {}
    cols = {c: excel_column(i) for i, c in enumerate(audit.columns)}
    for offset, record in enumerate(scored.to_dict("records"), start=2):
        ref = lambda title: f"{cols[title]}{offset}"
        def bind(expression: str) -> str:
            for title, column in cols.items():
                expression = expression.replace("{" + title + "}", f"{column}{offset}")
            if "{" in expression:
                raise ValueError(f"Unresolved rule formula placeholder: {expression}")
            return expression
        def formula(title: str, expression: str, expected: Any):
            address = ref(title)
            spec["formulas"][address] = "=" + expression
            spec["expected"][address] = _clean(expected)
        for r in RULE_CODES:
            condition, available, result = bind(conditions[r]), ref(f"{r} availability"), ref(f"{r} result")
            formula(f"{r} result", f'IF({condition},"FLAG",IF({available}="Not applicable","N/A",IF({available}="Not evaluable","NOT EVALUABLE","CLEAR")))',
                    {"Not applicable": "N/A", "Not evaluable": "NOT EVALUABLE"}.get(record[f"{r} status"], "FLAG" if record[r] else "CLEAR"))
            formula(f"{r} parity", f'IF(({result}="FLAG")=({ref(r + " Python flag")}=1),"PASS","ERROR")', "PASS")
        summands = "+".join(f'IF({ref(r + " result")}="FLAG",{rules[r]["points"]},0)' for r in RULE_CODES)
        formula("Score from Excel", summands, int(record["Risk score"]))
        formula("Score parity", f'IF({ref("Score from Excel")}={ref("Python risk score")},"PASS","ERROR")', "PASS")
        formula("Action from Excel", bind(conditions["decision"]), record["Recommended action"])
        formula("Decision parity", f'IF({ref("Action from Excel")}={ref("Python recommended action")},"PASS","ERROR")', "PASS")
    return spec


def build_export_payload(sources: Any, scored: pd.DataFrame) -> tuple[dict, pd.DataFrame]:
    from refined_screening import RULES
    registry = sources.registry
    public = _public_records(scored, RULES)
    if public[["PID", "Record ID"]].duplicated().any():
        raise ValueError("Duplicate PID and Record ID in final output")
    conditions, evidence = excel_audit_spec()
    sheets = [_frame_sheet("All Records", public)]
    for pid, study in registry.items():
        subset = public.loc[public["PID"].eq(str(pid))]
        sheet_name = study.get("sheet_name")
        if not sheet_name or len(sheet_name) > 31:
            raise ValueError(f"Invalid canonical Excel sheet name for PID {pid}")
        sheets.append(_frame_sheet(sheet_name, subset))
    exclusions = sources.exclusions.rename(columns={"project_id":"PID", "record_id":"Record ID", "response_key":"Response key", "exclusion_reason":"Exclusion reason", "matched_text_fields":"Fields containing test", "archive_clone":"Copied archive record", "prelaunch_protocol":"Prelaunch protocol record", "contains_test_word":"Contains the word test"}).copy()
    exclusions["PID"] = exclusions["PID"].astype(str)
    exclusions.insert(0,"Study",exclusions["PID"].map({str(p):s["display_label"] for p,s in registry.items()}))
    sheets.append(_frame_sheet("Excluded Records", exclusions))
    sheets.append(_rule_audit(scored, RULES, conditions, evidence))
    rules_frame = pd.DataFrame([{"Rule":r,"Check":v["label"],"Points":v["points"],"Evidence family":v.get("family",""),"Definition":v.get("description","")} for r,v in RULES.items()])
    notes = [
        {"cell":"A19","text":"FLAG means a screening check was triggered. ERROR means the Excel calculation disagrees with Python."},
        {"cell":"A21","text":"Unavailable or unasked questions are separate from clear answers. Identity and disability wording are review context, not proof of a bot."},
        {"cell":"A23","text":"Risk points summarize checks. Payment action also uses completion, contact readiness, and independent evidence families."},
        {"cell":"A25","text":"Excel recomputes thresholds and payment gates from visible normalized inputs. Metadata, branch visibility, language/choice decoding, postal format, text similarity and fingerprints are validated in Python."},
        {"cell":"A27","text":"University email or reference-cohort membership alone does not bypass screening. Decisions are snapshots: rerun the export after source corrections."},
        {"cell":"A29","text":"The literal free-text word test is a requested policy exclusion, including ordinary clinical discussion. Excluded records are not classified as bots."},
    ]
    sheets.append(_frame_sheet("Rule Definitions",rules_frame,notes=notes))
    provenance_names = {"project_id":"PID","study_name":"Study","api_project_title":"Live REDCap project title","population":"Population","source_mode":"Source mode","fetched_at_utc":"Retrieved at (UTC)","raw_records":"Source records","included_records":"Analyzed records","excluded_records":"Distinct excluded records","archive_clones_excluded":"Copied archive records","prelaunch_protocol_excluded":"Prelaunch protocol records","test_word_excluded":"Records containing test (may overlap)","source_mapping_review_records":"Records with mapping conflicts","record_sha256":"Record export SHA-256","metadata_sha256":"Metadata SHA-256"}
    sheets.append(_frame_sheet("Source Audit",sources.provenance.rename(columns=provenance_names)))
    field_labels = {"project_id":"PID", "source_field":"Source field", "canonical_field":"Canonical field",
                    "field_type":"Field type", "choice_codes":"Choice codes", "mapping_status":"Mapping status",
                    "invalid_response_count":"Invalid response count", "semantic_warning":"Question wording warning"}
    sheets.append(_frame_sheet("Field Mapping",sources.field_mapping.rename(columns=field_labels)))
    language_labels = {"project_id":"PID", "spanish_field":"Spanish field", "canonical_field":"Canonical field",
                       "answers_from_spanish":"Answers from Spanish", "conflicting_answers":"Conflicting answers",
                       "language_route_mismatches":"Language routing mismatches"}
    sheets.append(_frame_sheet("Language Audit",sources.folding_audit.rename(columns=language_labels)))
    summary_rows = []
    for pid, study in registry.items():
        subset = public.loc[public["PID"].eq(str(pid))]
        excluded = int(exclusions["PID"].eq(str(pid)).sum())
        source = sources.provenance.loc[sources.provenance["project_id"].astype(int).eq(int(pid))]
        if len(source) != 1:
            raise ValueError(f"Expected exactly one source provenance entry for PID {pid}")
        source = source.iloc[0]
        if int(source["included_records"]) != len(subset) or int(source["excluded_records"]) != excluded or int(source["raw_records"]) != len(subset)+excluded:
            raise ValueError(f"Source, analyzed and excluded record totals do not reconcile for PID {pid}")
        summary_rows.append({"Study":study["display_label"],"PID":str(pid),"Analyzed records":len(subset),**{a:int(subset["Recommended action"].eq(a).sum()) for a in ACTION_ORDER},"Excluded records":excluded,"Source records":int(source["raw_records"]),"Mapping review records":int(source["source_mapping_review_records"])})
    summary = pd.DataFrame(summary_rows)
    total = {c: int(summary[c].sum()) for c in summary.columns if c not in {"Study","PID"}}
    summary = pd.concat([summary,pd.DataFrame([{"Study":"All studies","PID":"All",**total}])],ignore_index=True)
    summary_spec = _frame_sheet("Summary",summary,headerRow=5,title="Caregiver response screening",notes=[
        {"cell":"A12","text":"Payment actions reflect survey completion, contact readiness, and corroborating evidence. The reason column explains each decision."},
        {"cell":"A14","text":"Exclusions remain visible in Excluded Records. The word test also excludes ordinary clinical discussion; these exclusions do not identify bots."},
        {"cell":"A16","text":"Study names follow the survey reference. Live REDCap identity, metadata, and record totals are documented in Source Audit."},
    ])
    summary_spec["formulas"],summary_spec["expected"] = {},{}
    end = len(public)+1
    for rowno,data in enumerate(summary.to_dict("records"),start=6):
        if data["PID"] == "All":
            for j,title in enumerate(summary.columns):
                if j < 2:
                    continue
                cell = f"{excel_column(j)}{rowno}"
                summary_spec["formulas"][cell] = f"=SUM({excel_column(j)}6:{excel_column(j)}{rowno-1})"
                summary_spec["expected"][cell] = data[title]
            continue
        formulas = {f"C{rowno}":f'=COUNTIFS(\'All Records\'!$B$2:$B${end},B{rowno})'}
        for j,action in enumerate(ACTION_ORDER,start=3):
            formulas[f"{excel_column(j)}{rowno}"] = f'=COUNTIFS(\'All Records\'!$B$2:$B${end},B{rowno},\'All Records\'!$D$2:$D${end},"{action}")'
        for cell,expression in formulas.items():
            summary_spec["formulas"][cell] = expression
            summary_spec["expected"][cell] = list(data.values())[ord(cell[0])-ord("A")]
    return {"sheets":[summary_spec,*sheets]},summary


def validate_workbook(path: Path, payload: dict, scored: pd.DataFrame) -> dict:
    """Read back actual formula caches; openpyxl is never used to author files."""
    import openpyxl
    values = openpyxl.load_workbook(path,data_only=True,read_only=False)
    formulas = openpyxl.load_workbook(path,data_only=False,read_only=False)
    expected_names = [s["name"] for s in payload["sheets"]]
    if values.sheetnames != expected_names:
        raise AssertionError("Workbook sheets differ from the validated export plan")
    count = 0
    for spec in payload["sheets"]:
        sheet = values[spec["name"]]
        for cell,expected in spec.get("expected",{}).items():
            if sheet[cell].value != expected:
                raise AssertionError(f"Cached formula value mismatch at {spec['name']}!{cell}")
            if formulas[spec["name"]][cell].data_type != "f":
                raise AssertionError(f"Audit is not an actual formula at {spec['name']}!{cell}")
            count += 1
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "e" or isinstance(cell.value,str) and cell.value in EXCEL_ERRORS | {"nan","<NA>","None","NaT"}:
                    raise AssertionError(f"Invalid workbook value at {spec['name']}!{cell.coordinate}")
    records = values["All Records"]
    keys = [(str(row[1]),str(row[2])) for row in records.iter_rows(min_row=2,values_only=True)]
    expected_keys = {(str(row["REDCap PID"]),str(row["Record ID"])) for row in scored.to_dict("records")}
    if len(keys) != len(set(keys)) or set(keys) != expected_keys:
        raise AssertionError("Saved workbook has duplicate, missing or extra study/record keys")
    values.close()
    formulas.close()
    return {"records":len(keys),"sheets":len(expected_names),"formula_checks":count,"formula_errors":0,"score_parity":"PASS","decision_parity":"PASS","rule_parity":"PASS","keys_unique":True}


def _artifact_modules() -> Path:
    candidates = [os.environ.get("CODEX_ARTIFACT_NODE_MODULES"),str(Path.home()/".cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules")]
    for candidate in candidates:
        if candidate and (Path(candidate)/"@oai/artifact-tool/package.json").is_file():
            return Path(candidate)
    raise RuntimeError("The Codex artifact-tool spreadsheet runtime is unavailable")


def write_refined_workbook(output_path: Path, payload: dict, scored: pd.DataFrame, *, render: bool = True) -> dict:
    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True,exist_ok=True)
    output_path.parent.chmod(0o700)
    native_engine = shutil.which("soffice") or shutil.which("libreoffice")
    if native_engine is None:
        raise RuntimeError("LibreOffice is required to verify saved Excel formula recalculation")
    with tempfile.TemporaryDirectory(prefix="caregiver_export_",dir=output_path.parent) as temporary:
        work = Path(temporary)
        (work/"node_modules").symlink_to(_artifact_modules(),target_is_directory=True)
        shutil.copy2(PROJECT_DIR/"refined_workbook.mjs",work/"builder.mjs")
        (work/"payload.json").write_text(json.dumps(payload,ensure_ascii=False,allow_nan=False),encoding="utf-8")
        author_file = work/"authored"/output_path.name
        command = [shutil.which("node") or "node",str(work/"builder.mjs"),str(work/"payload.json"),str(author_file)]
        if render:
            command.append(str(output_path.parent/"refined_validation_previews"))
        authored = subprocess.run(command,capture_output=True,text=True,timeout=600)
        if authored.returncode:
            raise RuntimeError("Spreadsheet authoring failed: " + authored.stderr[-3000:])
        authored_validation = validate_workbook(author_file,payload,scored)
        recalc_dir = work/"recalculated"
        recalc_dir.mkdir()
        profile = (work/"libreoffice-profile").as_uri()
        recalc = subprocess.run([native_engine,f"-env:UserInstallation={profile}","--headless","--convert-to","xlsx","--outdir",str(recalc_dir),str(author_file)],capture_output=True,text=True,timeout=600)
        final_file = recalc_dir/output_path.name
        if recalc.returncode or not final_file.exists():
            raise RuntimeError("Native Excel formula recalculation failed")
        validation = validate_workbook(final_file,payload,scored)
        validation.update({"authoring_engine":"artifact-tool","native_calculation_engine":"LibreOffice","native_recalculation":"PASS","artifact_formula_checks":authored_validation["formula_checks"]})
        if output_path.exists():
            backup = output_path.parent/"pre_refinement_backup"
            backup.mkdir(exist_ok=True)
            backup.chmod(0o700)
            stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
            shutil.copy2(output_path,backup/f"{output_path.stem}_{stamp}{output_path.suffix}")
        os.replace(final_file,output_path)
        output_path.chmod(0o600)
        validation_path = output_path.with_suffix(".validation.json")
        validation_path.write_text(json.dumps(validation,indent=2),encoding="utf-8")
        validation_path.chmod(0o600)
    return validation


def run_refined_export(project_dir: Path | str = PROJECT_DIR, refresh: bool = True,
                       output_path: Path | str | None = None) -> dict:
    from refined_sources import load_refined_sources
    from refined_screening import score_refined_records
    project_dir = Path(project_dir).resolve()
    output_path = Path(output_path) if output_path else project_dir/"Caregiver Outputs"/"restricted"/f"ESD_AllResponses_AllStudies_{date.today().isoformat()}.xlsx"
    sources = load_refined_sources(project_dir,refresh=refresh)
    scored = score_refined_records(sources)
    payload,summary = build_export_payload(sources,scored)
    validation = write_refined_workbook(output_path,payload,scored)
    return {"sources":sources,"scored":scored,"summary":summary,"validation":validation,"output_path":output_path.resolve()}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-dir",type=Path,default=PROJECT_DIR)
    parser.add_argument("--output",type=Path)
    parser.add_argument("--cached",action="store_true",help="Use only a previously identity-verified complete export cache")
    args = parser.parse_args()
    result = run_refined_export(args.project_dir,refresh=not args.cached,output_path=args.output)
    print(json.dumps({"output_path":str(result["output_path"]),**result["validation"]},indent=2))
